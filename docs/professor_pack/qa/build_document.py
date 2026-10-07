from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[3]
PACK=ROOT/'docs/professor_pack'
OUT=PACK/'qa'
font_path=Path('C:/Windows/Fonts/arial.ttf')
im=Image.new('RGB',(1600,335),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype(str(font_path),26)
small=ImageFont.truetype(str(font_path),22)
boxes=[(15,75,292,224),(332,75,609,224),(649,75,926,224),(966,75,1243,224),(1283,75,1580,224)]
labels=[('React client','Request and feedback'),('FastAPI','Validate and authenticate'),('Services','Candidates and scoring'),('SQL and artifacts','Profiles and vectors'),('Ranked JSON','Display recommendations')]
for box,(title,subtitle) in zip(boxes,labels):
    draw.rounded_rectangle(box,16,fill='#EEF3F8',outline='#8A99A8',width=2)
    cx=(box[0]+box[2])/2
    draw.text((cx,112),title,font=font,fill='black',anchor='mm')
    draw.text((cx,163),subtitle,font=small,fill='#293442',anchor='mm')
for i in range(4):
    x1=boxes[i][2]+7;x2=boxes[i+1][0]-7;y=148
    draw.line((x1,y,x2,y),fill='#526273',width=3)
    draw.polygon([(x2,y),(x2-9,y-6),(x2-9,y+6)],fill='#526273')
draw.text((800,285),'Offline scripts prepare features and models before online ranking',font=font,fill='black',anchor='mm')
diagram=OUT/'pipeline.png';im.save(diagram)

doc=Document();sec=doc.sections[0]
sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.3)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',23),('Heading 1',18),('Heading 2',12.5)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('PolyTaste  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(80,80,80)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def text_run(p,text):
    for i,part in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
        run=p.add_run(part[2:-2] if part.startswith('**') else part)
        if part.startswith('**'):run.bold=True

def table(lines):
    data=[[c.strip() for c in row.strip().strip('|').split('|')] for row in lines]
    data=[row for row in data if not all(re.fullmatch(r'[-: ]+',c) for c in row)]
    cols=len(data[0]);t=doc.add_table(rows=1,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if cols==4:widths=[3.1,.8,.8,2.0] if data[0][0]=='Variant' else [2.85,.5,2.85,.5]
    elif cols==5:widths=[2.25,.65,.65,.65,2.5]
    elif cols==3:widths=[2.2,1.45,3.05]
    else:widths=[6.7/cols]*cols
    for cell,width in zip(t.rows[0].cells,widths):cell.width=Inches(width)
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for ridx,row in enumerate(data):
        cells=t.rows[0].cells if ridx==0 else t.add_row().cells
        trpr=t.rows[ridx]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if ridx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cidx,(cell,content) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[cidx]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',65),('bottom',65),('left',85),('right',85)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243B53' if ridx==0 else ('F2F5F8' if ridx%2==0 else 'FFFFFF'));tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            if re.fullmatch(r'[\d., ]+(?:percent)?',content) or content in ['Not measured']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(content.replace('\\n','\n'));r.font.size=Pt(9 if cols>=4 else 9.5)
            if ridx==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

text=(PACK/'Project_Explanation.md').read_text(encoding='utf-8')
pages=text.split('--page--')
for page_no,page in enumerate(pages):
    lines=page.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(block);continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.0
                r=p.add_run(lines[i]);r.font.name='Consolas';r.font.size=Pt(9)
                i+=1
            i+=1;continue
        if line=='[[pipeline]]':doc.add_picture(str(diagram),width=Inches(6.8));i+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title' if page_no==0 else 'Heading 1')
            if page_no:p.paragraph_format.page_break_before=True
        elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 2')
        elif line.startswith('Source:'):
            p=doc.add_paragraph();r=p.add_run(line);r.italic=True;r.font.size=Pt(8)
            r.font.color.rgb=RGBColor(80,80,80)
        else:
            p=doc.add_paragraph();text_run(p,line)
        i+=1
doc.core_properties.title='PolyTaste Project Explanation and Viva Guide'
doc.core_properties.subject='Pipeline algorithms datasets evaluation and professor questions'
doc.core_properties.author=''
# The bundled default template can carry a blue title paragraph border.
for root in [doc.styles.element,doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
dest=PACK/'PolyTaste_Project_Explanation_and_Viva.docx';doc.save(dest)
# Structural QA supplements visual review; it does not replace rendering.
check=Document(dest)
assert len(check.tables)==10,len(check.tables)
assert len(pages)==19
assert any('32 What is your next improvement' in p.text for p in check.paragraphs)
print(json.dumps({'output':str(dest),'planned_pages':len(pages),'tables':len(check.tables),'paragraphs':len(check.paragraphs)}))
